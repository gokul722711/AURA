"""Deterministic offline tests for AURA document extraction and normalization."""

import io
from django.test import TestCase

from rag.exceptions import ExtractionError, UnsupportedFormatError
from rag.extraction import (
    DocxExtractor,
    ExtractedBlock,
    ExtractedDocument,
    MarkdownExtractor,
    PDFExtractor,
    TextExtractor,
    get_extractor,
    normalize_document,
    normalize_text,
)


def _make_pdf(pages: list[str]) -> bytes:
    """Create in-memory PDF bytes with given text per page."""
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        p = doc.new_page()
        if text.strip():
            p.insert_text((50, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx(
    elements: list[tuple[str, str, str | None]],  # (type, text, style)
    tables: list[list[list[str]]] | None = None,
) -> bytes:
    """Create in-memory DOCX bytes with given elements."""
    import docx

    doc = docx.Document()
    for elem_type, text, style in elements:
        if elem_type == "heading":
            doc.add_heading(text, level=1)
        elif elem_type == "list":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)

    if tables:
        for tbl_data in tables:
            t = doc.add_table(rows=len(tbl_data), cols=len(tbl_data[0]))
            for r_idx, row in enumerate(tbl_data):
                for c_idx, cell_text in enumerate(row):
                    t.rows[r_idx].cells[c_idx].text = cell_text

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TextExtractorTests(TestCase):
    """Tests for TextExtractor."""

    def setUp(self) -> None:
        self.extractor = TextExtractor()

    def test_extract_valid_text(self) -> None:
        content = b"Simple plain text content for testing."
        doc = self.extractor.extract(content, filename="test.txt")
        self.assertEqual(doc.source_type, "txt")
        self.assertIn("Simple plain text content", doc.text)
        self.assertEqual(len(doc.blocks), 1)
        self.assertEqual(doc.blocks[0].metadata["source_type"], "txt")

    def test_extract_empty_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"", filename="empty.txt")

    def test_extract_whitespace_only_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"   \n\n\t  ", filename="spaces.txt")


class MarkdownExtractorTests(TestCase):
    """Tests for MarkdownExtractor."""

    def setUp(self) -> None:
        self.extractor = MarkdownExtractor()

    def test_extract_valid_markdown(self) -> None:
        content = b"# AURA Architecture\n\n- Model Gateway\n- RAG Engine"
        doc = self.extractor.extract(content, filename="spec.md")
        self.assertEqual(doc.source_type, "markdown")
        self.assertIn("AURA Architecture", doc.text)
        self.assertEqual(len(doc.blocks), 1)
        self.assertEqual(doc.blocks[0].metadata["source_type"], "markdown")

    def test_extract_empty_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"", filename="empty.md")


class PDFExtractorTests(TestCase):
    """Tests for PDFExtractor using PyMuPDF."""

    def setUp(self) -> None:
        self.extractor = PDFExtractor()

    def test_extract_multipage_pdf_preserves_page_metadata(self) -> None:
        pdf_bytes = _make_pdf([
            "Page 1: Introduction to Quantum Mechanics.",
            "Page 2: Superposition and Entanglement principles.",
            "Page 3: Experimental Decoherence Benchmarks.",
        ])
        doc = self.extractor.extract(pdf_bytes, filename="quantum.pdf")

        self.assertEqual(doc.source_type, "pdf")
        self.assertEqual(doc.metadata["page_count"], 3)
        self.assertEqual(len(doc.blocks), 3)

        self.assertEqual(doc.blocks[0].metadata["page"], 1)
        self.assertEqual(doc.blocks[0].metadata["source_type"], "pdf")
        self.assertIn("Introduction to Quantum", doc.blocks[0].content)

        self.assertEqual(doc.blocks[1].metadata["page"], 2)
        self.assertIn("Superposition", doc.blocks[1].content)

        self.assertEqual(doc.blocks[2].metadata["page"], 3)
        self.assertIn("Decoherence", doc.blocks[2].content)

    def test_empty_pdf_bytes_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"", filename="zero.pdf")

    def test_corrupt_pdf_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"%PDF-1.4 corrupt junk data...", filename="corrupt.pdf")

    def test_pdf_with_no_text_fails(self) -> None:
        # 2 blank pages with no text (e.g. un-OCRed scanned PDF)
        pdf_bytes = _make_pdf(["", "   \n  "])
        with self.assertRaises(ExtractionError) as ctx:
            self.extractor.extract(pdf_bytes, filename="scanned.pdf")
        self.assertIn("contains no extractable text", str(ctx.exception).lower())


class DocxExtractorTests(TestCase):
    """Tests for DocxExtractor using python-docx."""

    def setUp(self) -> None:
        self.extractor = DocxExtractor()

    def test_extract_docx_structures(self) -> None:
        docx_bytes = _make_docx(
            elements=[
                ("heading", "Executive Summary", "Heading 1"),
                ("paragraph", "AURA coordinates autonomous reasoning agents.", None),
                ("list", "First milestone: Model Gateway", "List Bullet"),
                ("list", "Second milestone: Basic RAG", "List Bullet"),
            ],
            tables=[
                [["Parameter", "Value"], ["top_k", "5"], ["similarity", "0.85"]]
            ],
        )
        doc = self.extractor.extract(docx_bytes, filename="report.docx")

        self.assertEqual(doc.source_type, "docx")
        self.assertGreaterEqual(len(doc.blocks), 4)

        # Verify block types
        block_types = [b.metadata.get("block_type") for b in doc.blocks]
        self.assertIn("heading", block_types)
        self.assertIn("paragraph", block_types)
        self.assertIn("list", block_types)
        self.assertIn("table", block_types)

        # Verify table content
        table_block = next(b for b in doc.blocks if b.metadata.get("block_type") == "table")
        self.assertIn("Parameter | Value", table_block.content)
        self.assertIn("top_k | 5", table_block.content)

    def test_empty_docx_bytes_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"", filename="empty.docx")

    def test_corrupt_docx_fails(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extractor.extract(b"PK\x03\x04 not a real docx zip", filename="bad.docx")


class ExtractorRegistryTests(TestCase):
    """Tests for get_extractor factory and format validation."""

    def test_supported_extensions(self) -> None:
        self.assertIsInstance(get_extractor("file.txt"), TextExtractor)
        self.assertIsInstance(get_extractor("doc.md"), MarkdownExtractor)
        self.assertIsInstance(get_extractor("doc.markdown"), MarkdownExtractor)
        self.assertIsInstance(get_extractor("paper.pdf"), PDFExtractor)
        self.assertIsInstance(get_extractor("paper.docx"), DocxExtractor)

    def test_content_types(self) -> None:
        self.assertIsInstance(get_extractor(content_type="application/pdf"), PDFExtractor)
        self.assertIsInstance(get_extractor(content_type="text/plain"), TextExtractor)

    def test_unsupported_extension_fails(self) -> None:
        with self.assertRaises(UnsupportedFormatError) as ctx:
            get_extractor("image.png")
        self.assertIn(".png", str(ctx.exception))

        with self.assertRaises(UnsupportedFormatError):
            get_extractor("data.csv")

        with self.assertRaises(UnsupportedFormatError):
            get_extractor("binary.exe")


class NormalizationTests(TestCase):
    """Tests for normalization step."""

    def test_normalize_text_removes_artifacts(self) -> None:
        raw = "Line 1   \r\nLine 2\x00\x0c\r\n\n\n\n\nLine 3"
        cleaned = normalize_text(raw)
        self.assertNotIn("\x00", cleaned)
        self.assertNotIn("\x0c", cleaned)
        self.assertNotIn("\r", cleaned)
        self.assertIn("Line 1\nLine 2\n\nLine 3", cleaned)

    def test_normalize_document_drops_empty_blocks(self) -> None:
        doc = ExtractedDocument(
            text="Initial text",
            blocks=[
                ExtractedBlock(content="Real paragraph.", metadata={"page": 1}),
                ExtractedBlock(content="   \n\t  \x00", metadata={"page": 2}),
                ExtractedBlock(content="Second paragraph.", metadata={"page": 3}),
            ],
            source_type="pdf",
        )
        norm = normalize_document(doc)
        self.assertEqual(len(norm.blocks), 2)
        self.assertEqual(norm.blocks[0].metadata["page"], 1)
        self.assertEqual(norm.blocks[1].metadata["page"], 3)
        self.assertIn("Real paragraph.\n\nSecond paragraph.", norm.text)

    def test_normalize_empty_document_raises(self) -> None:
        doc = ExtractedDocument(
            text="",
            blocks=[ExtractedBlock(content="   \x00  ", metadata={})],
            source_type="txt",
        )
        with self.assertRaises(ExtractionError):
            normalize_document(doc)
