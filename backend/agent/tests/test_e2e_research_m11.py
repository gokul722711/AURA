"""End-to-end integration tests for M11 Knowledge Ingestion + Advanced Retrieval.

Verifies the complete flow:
  Upload/Ingest document (PDF / DOCX)
    ↓
  Knowledge Base
    ↓
  Knowledge Base research runtime
    ↓
  Retrieved evidence with page / block metadata
    ↓
  Grounded answer
    ↓
  Correct source and citation with page number
"""

import io
import json
from collections.abc import Iterator
from typing import Any

from django.test import TestCase
import pymupdf
import docx

from agent.research import create_research_runtime
from agent.results import make_citation
from agent.state import AgentStatus
from agent.tools.builtin.rag import RAGSearchTool
from agent.tools.policy import DefaultToolPolicy
from agent.tools.registry import ToolRegistry
from gateway.base import LLMProvider
from gateway.gateway import ModelGateway
from gateway.types import (
    ALL_CAPABILITIES,
    GenerationRequest,
    GenerationResponse,
    ProviderMetadata,
    StreamChunk,
    StructuredOutputRequest,
    StructuredOutputResponse,
)
from rag.chunking import ChunkingConfig
from rag.embeddings.mock import MockEmbeddingProvider
from rag.ingestion import ingest_file
from rag.retrieval import RetrievalConfig


class E2EScriptedLLMProvider(LLMProvider):
    """Deterministic LLM Provider returning a scripted sequence of responses."""

    def __init__(self, responses: list[str]) -> None:
        self.responses = list(responses)
        self.call_count = 0
        self.recorded_requests: list[GenerationRequest] = []
        self.recorded_structured_requests: list[StructuredOutputRequest] = []

    def metadata(self) -> ProviderMetadata:
        return ProviderMetadata(
            provider="scripted",
            model="scripted-model",
            capabilities=ALL_CAPABILITIES,
        )

    def generate(self, request: GenerationRequest) -> GenerationResponse:
        self.recorded_requests.append(request)
        self.call_count += 1
        if self.call_count <= len(self.responses):
            text = self.responses[self.call_count - 1]
        else:
            text = '{"decision": "finish"}'
        return GenerationResponse(
            text=text,
            provider="scripted",
            model="scripted-model",
        )

    def stream(self, request: GenerationRequest) -> Iterator[StreamChunk]:
        resp = self.generate(request)
        yield StreamChunk(text=resp.text, index=0, finish_reason="stop")

    def structured_output(
        self, request: StructuredOutputRequest
    ) -> StructuredOutputResponse:
        self.recorded_structured_requests.append(request)
        gen_req = GenerationRequest(
            messages=request.messages,
            model=request.model,
            temperature=request.temperature,
            max_tokens=request.max_tokens,
        )
        resp = self.generate(gen_req)
        try:
            data = json.loads(resp.text)
        except Exception:
            data = {"decision": "finish"}
        return StructuredOutputResponse(
            data=data,
            raw_text=resp.text,
            provider="scripted",
            model="scripted-model",
        )


def _build_pdf_bytes(pages_text: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _build_docx_bytes(paragraphs: list[tuple[str, str]]) -> bytes:
    doc = docx.Document()
    for text, style in paragraphs:
        if style.startswith("Heading"):
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class E2EResearchM11Tests(TestCase):
    """End-to-end tests for document ingestion, retrieval, evidence, and research citations."""

    def setUp(self) -> None:
        self.embedding_provider = MockEmbeddingProvider(dimensions=384)

    def test_pdf_research_flow_preserves_page_citations(self) -> None:
        """PDF ingestion -> RAG search -> Research runtime -> Citation with page number."""
        # 1. Build a 2-page PDF
        pdf_bytes = _build_pdf_bytes([
            "Overview of Neural Networks. Deep learning is a subset of machine learning.",
            "Attention Is All You Need. Transformer architectures rely on self-attention mechanisms to process tokens in parallel.",
        ])

        # 2. Ingest document
        doc = ingest_file(
            file_bytes=pdf_bytes,
            filename="transformer_paper.pdf",
            title="Transformer Architecture Research",
            embedding_provider=self.embedding_provider,
            chunking_config=ChunkingConfig(chunk_size=500, chunk_overlap=0),
            source="uploads/transformer_paper.pdf",
        )
        self.assertEqual(doc.status, "ready")
        self.assertEqual(doc.source_type, "pdf")

        # Find the chunk created for page 2
        chunks = list(doc.chunks.all().order_by("chunk_index"))
        self.assertGreaterEqual(len(chunks), 2)
        page2_chunk = None
        for chunk in chunks:
            if chunk.metadata.get("page") == 2:
                page2_chunk = chunk
                break
        self.assertIsNotNone(page2_chunk)
        chunk_id = str(page2_chunk.id)

        # Step 1: Decision to search knowledge base
        step1_decision = json.dumps({
            "decision": "continue",
            "query": "transformer self-attention",
        })
        # Step 2: Decision to finish research
        step2_decision = json.dumps({
            "decision": "finish",
        })
        # Step 3: Synthesis step generating grounded answer citing page 2
        canonical_citation = make_citation(doc.title, chunk_id, page=2)
        synthesis_response = (
            f"Transformers utilize self-attention mechanisms to process tokens concurrently {canonical_citation}."
        )

        llm_provider = E2EScriptedLLMProvider([
            step1_decision,
            step2_decision,
            synthesis_response,
        ])
        gateway = ModelGateway(provider=llm_provider)
        runtime = create_research_runtime(
            gateway=gateway,
            embedding_provider=self.embedding_provider,
            retrieval_config=RetrievalConfig(top_k=3),
            mode="knowledge_base",
        )

        # 4. Execute research
        result = runtime.run_research("Explain transformer self-attention architectures.")

        # 5. Verify research completion & status
        self.assertEqual(result.status, AgentStatus.COMPLETED)
        self.assertIn("self-attention mechanisms", result.final_answer)

        # 6. Verify evidence has page preserved
        self.assertGreater(len(result.evidence), 0)
        found_page2_evidence = False
        for ev in result.evidence:
            if ev.chunk_id == chunk_id:
                self.assertEqual(ev.page, 2)
                self.assertEqual(ev.metadata.get("source_type"), "pdf")
                found_page2_evidence = True
        self.assertTrue(found_page2_evidence, "Page 2 evidence was not collected")

        # 7. Verify citation integrity with page number preserved
        self.assertGreater(len(result.citations), 0)
        self.assertTrue(any("Page: 2" in c for c in result.citations))
        self.assertTrue(any("Page: 1" in c for c in result.citations))

        verification = result.verify_citations()
        self.assertIn(canonical_citation, verification["matched_citations"])
        self.assertIn(chunk_id, verification["matched_chunk_ids"])

        # 8. Verify source aggregation has pages
        self.assertGreater(len(result.sources), 0)
        source_entry = result.sources[0]
        self.assertEqual(source_entry["document_title"], "Transformer Architecture Research")
        self.assertIn(2, source_entry["pages"])

    def test_docx_research_flow_preserves_evidence_and_citations(self) -> None:
        """DOCX ingestion -> RAG search -> Research runtime -> Grounded answer with citations."""
        docx_bytes = _build_docx_bytes([
            ("Quantum Key Distribution Protocol", "Heading 1"),
            ("The BB84 protocol employs polarized photons to transmit quantum cryptographic keys securely.", "Normal"),
        ])

        doc = ingest_file(
            file_bytes=docx_bytes,
            filename="qkd_protocol.docx",
            title="QKD Protocol Specification",
            embedding_provider=self.embedding_provider,
            chunking_config=ChunkingConfig(chunk_size=500, chunk_overlap=0),
            source="uploads/qkd_protocol.docx",
        )
        self.assertEqual(doc.status, "ready")
        self.assertEqual(doc.source_type, "docx")

        chunks = list(doc.chunks.all())
        self.assertGreater(len(chunks), 0)
        chunk_id = str(chunks[0].id)

        canonical_citation = make_citation(doc.title, chunk_id)
        synthesis_response = (
            f"The BB84 protocol leverages polarized photons for quantum key transmission {canonical_citation}."
        )

        llm_provider = E2EScriptedLLMProvider([
            json.dumps({
                "decision": "continue",
                "query": "BB84 protocol polarized photons",
            }),
            json.dumps({"decision": "finish"}),
            synthesis_response,
        ])
        gateway = ModelGateway(provider=llm_provider)
        runtime = create_research_runtime(
            gateway=gateway,
            embedding_provider=self.embedding_provider,
            retrieval_config=RetrievalConfig(top_k=3),
            mode="knowledge_base",
        )

        result = runtime.run_research("How does BB84 protocol secure key transmission?")

        self.assertEqual(result.status, AgentStatus.COMPLETED)
        self.assertIn("polarized photons", result.final_answer)
        self.assertGreater(len(result.evidence), 0)
        self.assertEqual(result.evidence[0].metadata.get("source_type"), "docx")

        verification = result.verify_citations()
        self.assertIn(canonical_citation, verification["matched_citations"])
        self.assertIn(chunk_id, verification["matched_chunk_ids"])

    def test_document_ids_filtering_in_research_runtime(self) -> None:
        """Tool input document_ids constrains RAG retrieval during research."""
        doc_a = ingest_file(
            file_bytes=b"Alpha project documentation covers autonomous drone swarm algorithms.",
            filename="alpha.txt",
            title="Alpha Project",
            embedding_provider=self.embedding_provider,
            chunking_config=ChunkingConfig(chunk_size=500, chunk_overlap=0),
        )
        doc_b = ingest_file(
            file_bytes=b"Beta project documentation covers autonomous underwater submarine algorithms.",
            filename="beta.txt",
            title="Beta Project",
            embedding_provider=self.embedding_provider,
            chunking_config=ChunkingConfig(chunk_size=500, chunk_overlap=0),
        )

        chunk_a_id = str(doc_a.chunks.first().id)
        canonical_citation = make_citation(doc_a.title, chunk_a_id)

        # Agent specifically requests document_ids=[str(doc_a.id)]
        llm_provider = E2EScriptedLLMProvider([
            json.dumps({
                "decision": "continue",
                "query": "autonomous algorithms",
                "document_ids": [str(doc_a.id)],
            }),
            json.dumps({"decision": "finish"}),
            f"Alpha project focuses on autonomous drone swarms {canonical_citation}.",
        ])

        runtime = create_research_runtime(
            gateway=ModelGateway(provider=llm_provider),
            embedding_provider=self.embedding_provider,
            retrieval_config=RetrievalConfig(top_k=5),
            mode="knowledge_base",
        )

        result = runtime.run_research("Explain autonomous algorithms in Alpha project.")

        self.assertEqual(result.status, AgentStatus.COMPLETED)
        # All evidence gathered must only be from doc_a
        for ev in result.evidence:
            self.assertEqual(ev.document_id, str(doc_a.id))
            self.assertEqual(ev.document_title, "Alpha Project")
